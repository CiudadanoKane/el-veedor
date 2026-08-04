# Regenera src/lib/poblacionCiudades.js desde el anexo oficial del DANE.
#
# La tabla de población no se transcribe a mano: se lee del anexo de
# proyecciones municipales por área (base CNPV 2018), así que cada denominador
# per cápita se puede rastrear hasta la fuente primaria y rehacer cuando el DANE
# publique un año nuevo.
#
# Conserva la CURADURÍA (qué municipios componen el universo de comparación,
# ver el encabezado del módulo) y solo actualiza la cifra de cada uno. Para
# añadir o quitar municipios se edita ALTAS o el propio módulo, no este script.
#
# ABORTA si alguna fila curada no empareja con el anexo: perder un municipio en
# silencio lo sacaría del ranking sin que nadie lo notara.
#
# Uso:
#   pip install openpyxl
#   python scripts/generar-poblacion-municipal.py            # año por defecto
#   python scripts/generar-poblacion-municipal.py --anio 2026
import argparse, io, os, re, sys, unicodedata, urllib.request

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODULO = os.path.join(RAIZ, 'src', 'lib', 'poblacionCiudades.js')
ANEXO = os.path.join(RAIZ, '.cache-dane', 'PPED-AreaMun-2018-2042_VP.xlsx')
URL = ('https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/'
       'Municipal/PPED-AreaMun-2018-2042_VP.xlsx')

# Municipios donde SECOP usa el nombre corriente y el DANE el oficial. Se
# emparejan por CÓDIGO DANE, que es estable, en vez de por alias de nombre.
POR_CODIGO = {
    ('BOGOTA DC', 'BOGOTA'): '11001',            # DANE: "Bogotá. D.C."
    ('BOLIVAR', 'CARTAGENA'): '13001',           # DANE: "Cartagena de Indias"
    ('NORTE DE SANTANDER', 'CUCUTA'): '54001',   # DANE: "San José de Cúcuta"
    ('NARINO', 'TUMACO'): '52835',               # DANE: "San Andrés de Tumaco"
}

# Municipios a incluir que aún no estén en el módulo, por código DANE.
ALTAS = []

ALIAS_DEPTO = {
    'DISTRITO CAPITAL DE BOGOTA': 'BOGOTA DC',
    'SAN ANDRES PROVIDENCIA Y SANTA CATALINA':
        'ARCHIPIELAGO DE SAN ANDRES PROVIDENCIA Y SANTA CATALINA',
}


def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').upper())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Z ]', ' ', s)).strip()


def norm_depto(s):
    b = norm(s)
    return ALIAS_DEPTO.get(b, b)


def leer_anexo(anio):
    if not os.path.exists(ANEXO):
        os.makedirs(os.path.dirname(ANEXO), exist_ok=True)
        print(f'descargando el anexo del DANE…', file=sys.stderr)
        urllib.request.urlretrieve(URL, ANEXO)
    wb = openpyxl.load_workbook(ANEXO, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[2]]
    por_nombre, por_codigo = {}, {}
    for r in ws.iter_rows(min_row=10, values_only=True):
        if not r or r[4] is None or r[5] is None:
            continue
        if str(r[4]).strip() != anio or str(r[5]).strip().lower() != 'total':
            continue
        depto, cod, mpio, hab = str(r[1]).strip(), str(r[2]).strip(), str(r[3]).strip(), int(r[6])
        por_nombre[(norm_depto(depto), norm(mpio))] = hab
        por_codigo[cod] = (depto, mpio, hab)
    if not por_nombre:
        sys.exit(f'El anexo no trae datos para {anio}.')
    return por_nombre, por_codigo


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--anio', default='2025')
    args = ap.parse_args()

    por_nombre, por_codigo = leer_anexo(args.anio)
    print(f'anexo: {len(por_nombre)} municipios en {args.anio}', file=sys.stderr)

    src = io.open(MODULO, encoding='utf-8').read()
    ini = src.index('const POBLACION_MUNICIPAL_')
    ini = src.index('[', ini)
    fin = src.index('\n]', ini)
    filas = re.findall(r"\['([^']+)',\s*'([^']+)',\s*([\d_]+)\],([^\n]*)", src[ini:fin])
    if not filas:
        sys.exit('No se pudo leer la tabla actual del módulo.')

    salida, faltantes = [], []
    for depto, ciudad, _, nota in filas:
        k = (norm_depto(depto), norm(ciudad))
        hab = por_codigo[POR_CODIGO[k]][2] if k in POR_CODIGO else por_nombre.get(k)
        if hab is None:
            faltantes.append(f'{ciudad} ({depto})')
            continue
        salida.append((depto, ciudad, hab, nota.strip()))

    for cod, depto, ciudad, nota in ALTAS:
        if cod not in por_codigo:
            faltantes.append(f'{ciudad} ({depto}) por código {cod}')
            continue
        if any(norm(c) == norm(ciudad) and norm_depto(d) == norm_depto(depto)
               for d, c, _, _ in salida):
            continue
        salida.append((depto, ciudad, por_codigo[cod][2], nota))

    if faltantes:
        print('\nSIN EMPAREJAR (no se escribe nada):', file=sys.stderr)
        for f in faltantes:
            print('  -', f, file=sys.stderr)
        sys.exit(1)

    salida.sort(key=lambda r: -r[2])

    # El separador de miles se escribe con "_" (literal numérico de JavaScript).
    # OJO: la sustitución va SOLO sobre el número. Aplicarla a la línea entera
    # se lleva por delante las comas que separan los campos y deja el módulo
    # ilegible —pasó al escribir este script, y el propio generador lo detectó
    # en la corrida siguiente al no poder releer su salida—.
    def miles(n):
        return f'{n:,}'.replace(',', '_')

    tabla = '\n'.join(
        f"  ['{d}', '{c}', {miles(h)}],{(' ' + n) if n else ''}"
        for d, c, h, n in salida
    )
    nuevo = src[:ini + 1] + '\n' + tabla + src[fin:]
    nuevo = re.sub(r'POBLACION_CIUDADES_ANIO = \d+',
                   f'POBLACION_CIUDADES_ANIO = {args.anio}', nuevo)
    io.open(MODULO, 'w', encoding='utf-8', newline='\n').write(nuevo)
    print(f'escritas {len(salida)} filas en {MODULO}', file=sys.stderr)


if __name__ == '__main__':
    main()
